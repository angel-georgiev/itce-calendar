"""One-time import: Training_Schedule_2026.xlsx -> data/enrichment.json.

Keeps only the fields the website does not publish (instructor, category,
certification exam, relevant role) plus a corrected course link per code.
Everything else (dates, language, format, price) comes from itce.com via sync.py.

Usage: python3 scripts/import_xlsx.py path/to/Training_Schedule.xlsx
Requires openpyxl.
"""
import json
import re
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "data" / "enrichment.json"

CATEGORY_ALIASES = {
    "AI": "Artificial Intelligence (AI)",
}


def clean(v):
    if v is None:
        return ""
    return re.sub(r"\s+", " ", str(v)).strip()


def main(xlsx):
    ws = openpyxl.load_workbook(xlsx).active
    header = [clean(c.value) for c in ws[1]]
    col = {name: i for i, name in enumerate(header) if name}

    courses = {}
    instructors = {}
    rows = list(ws.iter_rows(min_row=2))
    # Oldest first, so later (newer) rows overwrite per-course values.
    rows.sort(key=lambda r: r[col["Start Date"]].value if isinstance(r[col["Start Date"]].value, datetime) else datetime.min)

    for r in rows:
        code = clean(r[col["Code"]].value)
        if not code:
            continue
        name_cell = r[col["Course name"]]
        c = courses.setdefault(code, {})
        c["title"] = clean(name_cell.value)

        category = clean(r[col["Training Category"]].value)
        category = CATEGORY_ALIASES.get(category, category)
        for key, value in (
            ("category", category),
            ("certification", clean(r[col["Certification Exam"]].value)),
            ("role", clean(r[col["Relevant Role"]].value)),
        ):
            if value:
                c[key] = value

        link_check = clean(r[col["Link Check"]].value)
        m = re.search(r"page: (https?://\S+)", link_check)
        if m:
            c["link"] = m.group(1)
        elif link_check.startswith("OK") and name_cell.hyperlink and name_cell.hyperlink.target:
            c["link"] = name_cell.hyperlink.target

        start = r[col["Start Date"]].value
        instructor = clean(r[col["Primary Instructor"]].value)
        if isinstance(start, datetime) and instructor:
            instructors[f"{code}|{start:%Y-%m-%d}"] = instructor

    data = {
        "_note": "Fields not published on itce.com. Edit by hand or via chat; sync.py merges them by course code.",
        "courses": dict(sorted(courses.items())),
        "instructors": dict(sorted(instructors.items())),
    }
    OUT.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Wrote {OUT}: {len(courses)} courses, {len(instructors)} instructor assignments")


if __name__ == "__main__":
    main(sys.argv[1])
